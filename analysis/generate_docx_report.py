import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_report():
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    NAVY = RGBColor(27, 54, 93)       # Primary
    SLATE = RGBColor(92, 118, 141)    # Secondary
    CORAL = RGBColor(217, 83, 79)     # OOM Alert

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = title.add_run("EMPIRICAL BASELINE PROFILING REPORT")
    r_title.bold = True
    r_title.font.size = Pt(20)
    r_title.font.name = "Segoe UI"
    r_title.font.color.rgb = NAVY

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = subtitle.add_run("Multi-Model KV-Cache Scaling & Dual OOM Ceiling Mapping (8GB & 16GB Targets)")
    r_sub.font.size = Pt(12.5)
    r_sub.font.name = "Segoe UI Semibold"
    r_sub.font.color.rgb = SLATE

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = meta.add_run("CPE401 Senior Project I · Computer Engineering, KMUTT\nTestbed: NVIDIA RTX 4090 (24GB) @ CPE Slurm Cluster · Date: 19 Oct 2026")
    r_meta.font.size = Pt(9.5)
    r_meta.font.name = "Segoe UI"
    r_meta.font.color.rgb = SLATE

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 1. Executive Summary
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Executive Summary & Dual OOM Breakthroughs")
    r_h1.font.color.rgb = NAVY
    r_h1.font.name = "Segoe UI"

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.line_spacing = 1.15
    p_intro.add_run(
        "This empirical report documents native FP16/BF16 autoregressive LLM inference benchmarks across four representative open models "
        "(Qwen2.5-1.5B, Qwen2.5-3B, Phi-3.5-mini 3.8B, and Qwen2.5-7B) on the KMUTT CPE Slurm Cluster. "
        "The findings prove the existence of dual hardware bottlenecks outlined in the project proposal:"
    )

    bullets = [
        ("The 8GB Consumer OOM Ceiling (Qwen2.5-3B): ", "At 32K context, total peak VRAM hits 9,747.98 MB, exceeding the 8GB limit (8,192 MB) by +1,555.98 MB (+19.0%), guaranteeing a fatal CUDA OOM on local RTX 4060 GPUs."),
        ("The 16GB Cluster/Cloud OOM Ceiling (Qwen2.5-7B): ", "Static model weights consume 14,568.65 MB (~14.23 GB). At 16K context, VRAM reaches 17,705.46 MB (+1,321 MB over 16GB), and at 32K context reaches 20,834.15 MB (+4,450 MB / +27.2% over 16GB)."),
        ("Cross-Architecture MHA Vulnerability (Phi-3.5-mini 3.8B): ", "Due to eager attention and standard MHA overhead, Phi-3.5 breaches the 8GB limit at just 1,024 tokens (8,410.93 MB) and crashes with CUDA OOM at 8K tokens even on a 24GB RTX 4090."),
        ("Severe Latency Degradation (TPOT Inflation): ", "At 32K context, Qwen2.5-7B token decoding latency slows down to 130.39 ms/token (a 6.6× slowdown compared to 19.81 ms at 1K tokens), illustrating heavy memory bandwidth saturation.")
    ]
    for b_title, b_desc in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r_bt = bp.add_run(b_title)
        r_bt.bold = True
        r_bt.font.color.rgb = NAVY
        bp.add_run(b_desc)

    # 2. Testbed Specs Table
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. Hardware & Software Telemetry")
    r_h2.font.color.rgb = NAVY
    r_h2.font.name = "Segoe UI"

    spec_table = doc.add_table(rows=5, cols=2)
    spec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    spec_data = [
        ("Compute Host", "gpu01.slurm.cpe.kmutt.ac.th (KMUTT CPE Slurm Cluster)"),
        ("GPU Hardware", "NVIDIA GeForce RTX 4090 (24,080.88 MB Total VRAM)"),
        ("Driver & CUDA Stack", "Driver 580.178.04 | CUDA 13.0 (PyTorch CUDA 12.1 runtime)"),
        ("Execution Environment", "Python 3.12 | PyTorch 2.5.1+cu121 | Transformers 4.47.1 | nvidia-ml-py"),
        ("Slurm Execution", "Job IDs: 22091 (Qwen 1.5B/3B) & 22094 (Phi-3.5 3.8B & Qwen 7B)")
    ]
    for i, (k, v) in enumerate(spec_data):
        row = spec_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(2.2), Inches(4.8)
        c0.text, c1.text = k, v
        c0.paragraphs[0].runs[0].bold = True
        set_cell_background(c0, "F0F4F8")
        set_cell_margins(c0)
        set_cell_margins(c1)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Table Helper
    def add_benchmark_table(heading_text, sub_text, headers, rows, oom_threshold_str=None):
        h = doc.add_heading(level=1)
        rh = h.add_run(heading_text)
        rh.font.color.rgb = NAVY
        rh.font.name = "Segoe UI"
        
        p = doc.add_paragraph()
        p.add_run(sub_text).font.color.rgb = SLATE
        p.paragraph_format.space_after = Pt(4)

        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for col_idx, h_text in enumerate(headers):
            cell = table.rows[0].cells[col_idx]
            cell.text = h_text
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            set_cell_background(cell, "1B365D")
            set_cell_margins(cell, 80, 80, 90, 90)

        for row_idx, data in enumerate(rows):
            row = table.rows[row_idx + 1]
            is_oom = "OOM" in data[-1]
            bg = "FDF2F2" if is_oom else ("FFFFFF" if row_idx % 2 == 0 else "F8FAFC")
            for col_idx, val in enumerate(data):
                cell = row.cells[col_idx]
                cell.text = val
                if is_oom and col_idx == len(data) - 1:
                    cell.paragraphs[0].runs[0].bold = True
                    cell.paragraphs[0].runs[0].font.color.rgb = CORAL
                set_cell_background(cell, bg)
                set_cell_margins(cell, 60, 60, 80, 80)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 3. Tier 1: Qwen2.5-1.5B
    add_benchmark_table(
        "3. Tier 1: Qwen2.5-1.5B-Instruct (Consumer 8GB Small Model)",
        "Static Model Weights: 2,945.29 MB (~2.88 GB) | Headroom on 8GB GPU: 5,246.71 MB",
        ["Context", "Peak VRAM (MB)", "Est. KV-Cache (MB)", "KV Share", "TPOT (ms/token)", "Status on 8GB GPU (RTX 4060)"],
        [
            ("512", "3,000.94", "55.65", "1.9%", "16.67", "PASS (36.6% VRAM)"),
            ("1,024", "3,046.96", "101.67", "3.3%", "8.69", "PASS (37.2% VRAM)"),
            ("2,048", "3,142.50", "197.21", "6.3%", "9.07", "PASS (38.4% VRAM)"),
            ("4,096", "3,325.59", "380.30", "11.4%", "9.85", "PASS (40.6% VRAM)"),
            ("8,192", "3,697.76", "752.47", "20.4%", "12.29", "PASS (45.1% VRAM)"),
            ("16,384", "4,442.11", "1,496.81", "33.7%", "21.29", "PASS (54.2% VRAM)"),
            ("32,768", "5,930.79", "2,985.50", "50.3%", "44.24", "PASS (72.4% VRAM - TPOT 5.1x slow)")
        ]
    )

    # 4. Tier 2: Qwen2.5-3B
    add_benchmark_table(
        "4. Tier 2: Qwen2.5-3B-Instruct (Consumer 8GB OOM Proof)",
        "Static Model Weights: 5,994.48 MB (~5.85 GB) | Headroom on 8GB GPU: Only 2,197.52 MB",
        ["Context", "Peak VRAM (MB)", "Est. KV-Cache (MB)", "KV Share", "TPOT (ms/token)", "Status on 8GB GPU (RTX 4060)"],
        [
            ("512", "6,061.13", "66.65", "1.1%", "13.56", "PASS (74.0% VRAM)"),
            ("1,024", "6,121.15", "126.67", "2.1%", "11.44", "PASS (74.7% VRAM)"),
            ("2,048", "6,239.69", "245.21", "3.9%", "12.69", "PASS (76.2% VRAM)"),
            ("4,096", "6,470.78", "476.30", "7.4%", "15.27", "PASS (79.0% VRAM)"),
            ("8,192", "6,938.95", "944.47", "13.6%", "21.31", "High Stress (84.7% VRAM)"),
            ("16,384", "7,875.29", "1,880.81", "23.9%", "38.64", "Danger Zone (96.1% VRAM)"),
            ("32,768", "9,747.98", "3,753.50", "38.5%", "76.74", "FATAL OOM CRASH (+1.55 GB)")
        ]
    )

    # 5. Tier 3: Phi-3.5-mini 3.8B
    add_benchmark_table(
        "5. Tier 3: Phi-3.5-mini-instruct (3.8B Cross-Architecture Comparison)",
        "Static Model Weights: 7,288.38 MB (~7.12 GB) | Headroom on 8GB GPU: Only 903.62 MB",
        ["Context", "Peak VRAM (MB)", "Est. KV-Cache (MB)", "TPOT (ms/token)", "Status on 8GB GPU", "Status on 24GB RTX 4090"],
        [
            ("512", "7,780.16", "491.78", "34.46", "Near Limit (95.0%)", "PASS"),
            ("1,024", "8,410.93", "1,122.55", "22.28", "FATAL OOM CRASH (1K)", "PASS"),
            ("2,048", "10,169.34", "2,880.96", "25.23", "FATAL OOM CRASH", "PASS"),
            ("4,096", "18,695.64", "11,407.26", "51.04", "FATAL OOM CRASH", "Heavy Load (77.6%)"),
            ("8,192", "-", "-", "-", "FATAL OOM CRASH", "FATAL OOM CRASH (24GB)")
        ]
    )

    # 6. Tier 4: Qwen2.5-7B
    add_benchmark_table(
        "6. Tier 4: Qwen2.5-7B-Instruct (Cluster/Cloud 16GB Tier)",
        "Static Model Weights: 14,568.65 MB (~14.23 GB) | Headroom on 16GB GPU: Only 1,815.35 MB",
        ["Context", "Peak VRAM (MB)", "Est. KV-Cache (MB)", "KV Share", "TPOT (ms/token)", "Status on 16GB Cluster GPU"],
        [
            ("512", "14,674.55", "105.90", "0.7%", "28.49", "PASS (89.6% VRAM)"),
            ("1,024", "14,775.32", "206.67", "1.4%", "19.81", "PASS (90.2% VRAM)"),
            ("2,048", "14,974.24", "405.59", "2.7%", "22.08", "PASS (91.4% VRAM)"),
            ("4,096", "15,359.82", "791.17", "5.2%", "27.52", "High Stress (93.7% VRAM)"),
            ("8,192", "16,142.12", "1,573.47", "9.8%", "39.97", "Danger Zone (98.5% VRAM)"),
            ("16,384", "17,705.46", "3,136.81", "17.7%", "67.07", "FATAL OOM CRASH (+1.32 GB)"),
            ("32,768", "20,834.15", "6,265.50", "30.1%", "130.39", "FATAL OOM CRASH (+4.45 GB)")
        ]
    )

    # 7. Research Significance
    h7 = doc.add_heading(level=1)
    rh7 = h7.add_run("7. Research Significance & Proposal Alignment")
    rh7.font.color.rgb = NAVY
    rh7.font.name = "Segoe UI"

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.line_spacing = 1.15
    p_sig.add_run(
        "These comprehensive results firmly establish the necessity and commercial viability of the proposed Co-Optimization system:\n"
        "1. Consumer Local GPUs (8GB): Selective 4-bit compression is mandatory to compress the 32K KV-Cache of 3B models from 3.75 GB down to ~1.1 GB, preventing fatal OOM.\n"
        "2. Cluster/Cloud GPUs (16GB): For 7B models, 32K KV-Cache (6.27 GB) causes total VRAM to reach 20.83 GB (+27.2% above limit). "
        "Asynchronous PCIe host RAM offload combined with 4-bit compression will compress and page out older context tokens, enabling full 32K execution on a single 16GB instance without crashing."
    )

    out_path = os.path.join(os.path.dirname(__file__), "BASELINE_PROFILING_REPORT_RTX4090.docx")
    doc.save(out_path)
    print(f"[SUCCESS] Multi-Model Word report generated: {out_path}")

if __name__ == "__main__":
    create_report()
